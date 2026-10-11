
from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_COLLISION_PATTERN = re.compile(
    r"tanımı.*ölçü.*birimi.*miktarı.*birim\s*fiyatı",
    re.IGNORECASE | re.DOTALL,
)

POZ_PATTERN = re.compile(r"(?<!\d)\d{2}\.\d{3}\.\d{4}(?!\d)")


def parse_turkish_number(value: Any) -> float | None:
    """Convert a Turkish-formatted number to float; return None if uncertain."""
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("\u00a0", "").replace(" ", "")
    text = text.replace("₺", "").replace("TL", "").replace("tl", "")

    # A slash can represent a ratio or formula, not a single numeric value.
    if "/" in text:
        return None

    # Remove thousands separators and convert the decimal comma.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    text = text.replace("%", "")
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def sha256_file(path: str | Path) -> str:
    """Return the SHA-256 digest of a file's actual bytes."""
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_suspicious_header(text: str) -> bool:
    normalized = " ".join((text or "").split())
    return bool(HEADER_COLLISION_PATTERN.search(normalized))


def extract_pdf(pdf_path: str | Path) -> dict[str, Any]:
    """Extract page text and tables while preserving page/table references."""
    source = Path(pdf_path)
    if not source.is_file():
        raise FileNotFoundError(f"PDF bulunamadı: {source}")
    if source.suffix.lower() != ".pdf":
        raise ValueError("Girdi dosyası PDF olmalıdır.")

    pages: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []

    with pdfplumber.open(source) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            pages.append({"page": page_number, "text": page_text})

            extracted_tables = page.extract_tables() or []
            for table_number, table in enumerate(extracted_tables, start=1):
                for row_number, row in enumerate(table, start=1):
                    cells = ["" if cell is None else str(cell).strip()
                             for cell in row]
                    row_text = " | ".join(cells)

                    record = {
                        "page": page_number,
                        "table": table_number,
                        "row": row_number,
                        "cells": cells,
                        "text": row_text,
                    }
                    tables.append(record)

                    if _is_suspicious_header(row_text):
                        warnings.append({
                            "severity": "HIGH",
                            "page": page_number,
                            "table": table_number,
                            "row": row_number,
                            "reason": (
                                "Başlık metni veri satırına karışmış olabilir."
                            ),
                            "source_text": row_text,
                            "status": "REVIEW_REQUIRED",
                        })

            # Also flag the known header-collision phrase if it appears in text.
            for line_number, line in enumerate(page_text.splitlines(), start=1):
                if _is_suspicious_header(line):
                    warnings.append({
                        "severity": "HIGH",
                        "page": page_number,
                        "table": None,
                        "row": line_number,
                        "reason": (
                            "Sayfa metninde başlık/veri çakışması şüphesi."
                        ),
                        "source_text": line,
                        "status": "REVIEW_REQUIRED",
                    })

    return {
        "source_file": str(source),
        "source_sha256": sha256_file(source),
        "pages": pages,
        "tables": tables,
        "warnings": warnings,
    }


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(
            (len(str(cell.value or "")) for cell in column_cells),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(max_length + 2, 12), 55)

    # Highlight rows that explicitly require review.
    if ws.title == "Kontrol ve Onay":
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.fill = warning_fill


def write_excel(data: dict[str, Any], output_path: str | Path) -> Path:
    """Write extracted tables and review warnings to a four-sheet workbook."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    poz_ws = workbook.active
    poz_ws.title = "Poz Analizleri"
    detail_ws = workbook.create_sheet("Detay Kalemler")
    review_ws = workbook.create_sheet("Kontrol ve Onay")
    summary_ws = workbook.create_sheet("Test Özeti")

    poz_ws.append([
        "Kaynak Sayfa", "Tablo", "Satır", "Poz No", "Kaynak Satır"
    ])
    detail_ws.append([
        "Kaynak Sayfa", "Tablo", "Satır", "Sütun", "Ham Değer",
        "Sayısal Değer"
    ])
    review_ws.append([
        "Önem", "Kaynak Sayfa", "Tablo", "Satır", "Açıklama",
        "Kaynak Metin", "Durum"
    ])
    summary_ws.append(["Kontrol", "Sonuç"])

    for table in data["tables"]:
        row_text = table["text"]
        poz_matches = POZ_PATTERN.findall(row_text)

        for poz_no in poz_matches:
            poz_ws.append([
                table["page"], table["table"], table["row"],
                poz_no, row_text
            ])

        for column_number, raw_value in enumerate(table["cells"], start=1):
            detail_ws.append([
                table["page"],
                table["table"],
                table["row"],
                column_number,
                raw_value,
                parse_turkish_number(raw_value),
            ])

    for warning in data["warnings"]:
        review_ws.append([
            warning["severity"],
            warning["page"],
            warning["table"],
            warning["row"],
            warning["reason"],
            warning["source_text"],
            warning["status"],
        ])

    summary_ws.append(["Kaynak PDF", data["source_file"]])
    summary_ws.append(["SHA-256", data["source_sha256"]])
    summary_ws.append(["Sayfa sayısı", len(data["pages"])])
    summary_ws.append(["Çıkarılan tablo satırı", len(data["tables"])])
    summary_ws.append(["Bulunan poz eşleşmesi", poz_ws.max_row - 1])
    summary_ws.append(["İnceleme uyarısı", len(data["warnings"])])
    summary_ws.append([
        "Doğrulama durumu",
        "Taslak çıktı; kaynakla karşılaştırma ve insan onayı gerekir."
    ])

    for ws in workbook.worksheets:
        _style_sheet(ws)

    workbook.save(output)
    return output


def write_word(data: dict[str, Any], output_path: str | Path) -> Path:
    """Create a review report with source metadata, warnings and extracted text."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("CHATIN — PDF İnceleme Raporu", level=0)
    doc.add_paragraph(f"Kaynak dosya: {data['source_file']}")
    doc.add_paragraph(f"SHA-256: {data['source_sha256']}")
    doc.add_paragraph(
        "Bu rapor otomatik çıkarımdır. Kaynağın doğruluğunu tek başına "
        "kanıtlamaz; şüpheli kayıtlar kullanıcı tarafından incelenmelidir."
    )

    doc.add_heading("1. Kontrol ve Onay", level=1)
    if not data["warnings"]:
        doc.add_paragraph("Otomatik kontrollerde tanımlı şüpheli ifade bulunmadı.")
        doc.add_paragraph(
            "Bu, PDF'nin hatasız olduğu anlamına gelmez."
        )
    else:
        for warning in data["warnings"]:
            paragraph = doc.add_paragraph()
            paragraph.add_run(
                f"[{warning['severity']}] Sayfa {warning['page']}, "
                f"tablo {warning['table']}, satır {warning['row']}: "
            ).bold = True
            paragraph.add_run(warning["reason"])
            paragraph = doc.add_paragraph(warning["source_text"])
            for run in paragraph.runs:
                run.font.highlight_color = 7  # Word yellow highlight

    doc.add_heading("2. Çıkarılan Tablolar", level=1)
    for table in data["tables"]:
        doc.add_heading(
            f"Sayfa {table['page']} — Tablo {table['table']} — "
            f"Satır {table['row']}",
            level=2,
        )
        doc.add_paragraph(table["text"] or "(Boş satır)")

    doc.add_heading("3. Sayfa Metinleri", level=1)
    for page in data["pages"]:
        doc.add_heading(f"Sayfa {page['page']}", level=2)
        doc.add_paragraph(page["text"] or "(Metin çıkarılamadı)")

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    doc.save(output)
    return output


class ChatinDocumentEngine:
    """Small initial PDF-to-Excel/Word extraction engine."""

    sha256_file = staticmethod(sha256_file)
    parse_turkish_number = staticmethod(parse_turkish_number)
    extract_pdf = staticmethod(extract_pdf)
    write_excel = staticmethod(write_excel)
    write_word = staticmethod(write_word)

    def process_pdf(
        self,
        pdf_path: str | Path,
        output_dir: str | Path = "output",
    ) -> dict[str, str]:
        data = extract_pdf(pdf_path)
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        stem = Path(pdf_path).stem
        excel_path = write_excel(data, output_dir / f"{stem}_CHATIN.xlsx")
        word_path = write_word(data, output_dir / f"{stem}_CHATIN.docx")

        return {
            "excel": str(excel_path),
            "word": str(word_path),
            "source_sha256": data["source_sha256"],
            "warnings": str(len(data["warnings"])),
        }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="CHATIN PDF → Excel ve Word başlangıç motoru"
    )
    parser.add_argument("pdf", help="İşlenecek PDF dosyası")
    parser.add_argument(
        "--output-dir", default="output",
        help="Çıktı klasörü (varsayılan: output)",
    )
    args = parser.parse_args()

    engine = ChatinDocumentEngine()
    results = engine.process_pdf(args.pdf, args.output_dir)

    print("İşlem tamamlandı.")
    print(f"Excel: {results['excel']}")
    print(f"Word: {results['word']}")
    print(f"Kaynak SHA-256: {results['source_sha256']}")
    print(f"İnceleme uyarısı: {results['warnings']}")
    print("Not: Dosyalar taslaktır; kaynakla karşılaştırılmalıdır.")


if __name__ == "__main__":
    main()
