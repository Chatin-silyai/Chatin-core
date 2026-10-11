
import hashlib

from openpyxl import load_workbook
from docx import Document

from app import (
    ChatinDocumentEngine,
    parse_turkish_number,
    write_excel,
    write_word,
)


def test_parse_turkish_number():
    assert parse_turkish_number("1.234,56") == 1234.56
    assert parse_turkish_number("123,45") == 123.45
    assert parse_turkish_number("100") == 100.0
    assert parse_turkish_number("0,085/1,80") is None
    assert parse_turkish_number("belirsiz") is None
    assert parse_turkish_number(None) is None


def test_sha256_file(tmp_path):
    source = tmp_path / "sample.txt"
    content = b"CHATIN test"
    source.write_bytes(content)

    expected = hashlib.sha256(content).hexdigest()
    assert ChatinDocumentEngine.sha256_file(source) == expected


def sample_data():
    return {
        "source_file": "ornek.pdf",
        "source_sha256": "a" * 64,
        "pages": [{"page": 1, "text": "Örnek kaynak metni"}],
        "tables": [
            {
                "page": 1,
                "table": 1,
                "row": 1,
                "cells": ["19.100.1121", "1.234,50"],
                "text": "19.100.1121 | 1.234,50",
            }
        ],
        "warnings": [
            {
                "severity": "HIGH",
                "page": 1,
                "table": 1,
                "row": 2,
                "reason": "Başlık/veri çakışması şüphesi.",
                "source_text": (
                    "Tanımı Ölçü Birimi Miktarı Birim Fiyatı"
                ),
                "status": "REVIEW_REQUIRED",
            }
        ],
    }


def test_write_excel_creates_expected_sheets(tmp_path):
    output = tmp_path / "result.xlsx"
    write_excel(sample_data(), output)

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == [
        "Poz Analizleri",
        "Detay Kalemler",
        "Kontrol ve Onay",
        "Test Özeti",
    ]
    assert workbook["Poz Analizleri"].max_row == 2
    assert workbook["Kontrol ve Onay"].max_row == 2


def test_write_word_creates_report(tmp_path):
    output = tmp_path / "result.docx"
    write_word(sample_data(), output)

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "CHATIN" in text
    assert "Başlık/veri çakışması şüphesi." in text
    assert "Örnek kaynak metni" in text
